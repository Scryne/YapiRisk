"""Yöntemin ölçülmüş isabeti (K-6): ACI-133 (2023 Kahramanmaraş, 242 BA bina, CC BY 4.0).

Yöntem ve isabet sayfasının sayıları buradan gelir. Ham veri yayımlanmaz; yalnız özet istatistik.
Veride olmayan RYTEİE girdileri (görünen kalite, ağır çıkma, düşeyde/planda düzensizlik, nizam, eğim)
"yok" sayılır: ölçülen şey kısmi puandır. Örneklem hasar keşfi için seçilmiştir; AUC mutlak isabet değil,
sıralama gücüdür.
"""

from openpyxl import load_workbook

from . import yontem
from .ortak import indir
from .tehlike import bilineer, izgara_cek

ACI133 = "https://zenodo.org/api/records/13386343/files/ACI133v3.xlsx/content"
ACI133_KUNYE = {
    "ad": "ACI 133 — Damage to RC Structures and Intensity Measures from the 2023 Türkiye Earthquake",
    "yazar": "Hariri-Ardebili 2024",
    "url": "https://zenodo.org/records/13386343",
    "lisans": "CC BY 4.0",
}
#: S1: ESHM20/TDTH SS oranı; ürünün kötü ucu SS/0,72. Doğrulama üç varsayımı da raporlar.
ORANLAR = {"ham": 1.0, "ortalama": 0.89, "kotu": 0.72}
BANTLAR = [("yuksek", float("-inf"), 50), ("orta", 50, 100), ("dusuk", 100, float("inf"))]
OLCULEMEYEN = [
    "gorunenKalite",
    "agirCikma",
    "duseydeDuzensizlik",
    "plandaDuzensizlik",
    "yapiNizami",
    "zeminEgimi",
]


def auc(riskli_skor: list[float], hedef: list[bool]) -> float:
    """Mann–Whitney AUC; skor büyükse daha riskli, eşitlik yarım."""
    poz = [s for s, t in zip(riskli_skor, hedef, strict=True) if t]
    neg = [s for s, t in zip(riskli_skor, hedef, strict=True) if not t]
    toplam = sum(1.0 if p > n else 0.5 if p == n else 0.0 for p in poz for n in neg)
    return toplam / (len(poz) * len(neg))


def binalari_oku(yol) -> list[dict]:
    """ACI133v3.xlsx: ilk 3 satır başlık; sütunlar spike/S3/aci133_analiz.py ile aynı."""
    ws = load_workbook(yol, read_only=True, data_only=True).worksheets[0]
    binalar = []
    for s in list(ws.iter_rows(values_only=True))[3:]:
        if s[0] is None:
            continue
        binalar.append(
            {
                "yas": float(s[1]),
                "kat": int(s[2]),
                "perde_alan": float(s[14]),
                "hasar3": s[21],
                "kritik": s[23],
                "kisa_kolon": s[27] == "Y",
                "pi": float(s[33]),
                "yumusak_kat": s[34] == "Y",
                "enlem": float(s[36]),
                "boylam": float(s[37]),
                "vs30": float(s[38]),
            }
        )
    return binalar


def olc() -> tuple[dict, dict]:
    xlsx, xlsx_meta = indir("dogrulama/ACI133v3.xlsx", ACI133)
    tum = binalari_oku(xlsx)
    kutu = (
        min(b["boylam"] for b in tum),
        min(b["enlem"] for b in tum),
        max(b["boylam"] for b in tum),
        max(b["enlem"] for b in tum),
    )
    izgaralar, izgara_meta = izgara_cek("dogrulama/eshm20", kutu, eksik_serbest=True)
    kapsam = [b for b in tum if 1 <= b["kat"] <= 7]
    hedef = [b["hasar3"] == "S" for b in kapsam]

    def pp(b: dict, oran: float, perde: bool) -> int:
        ss = bilineer(izgaralar["SS"], b["enlem"], b["boylam"]) / oran
        zemin = yontem.zemin_sinifi(b["vs30"])
        bolge = yontem.tehlike_bolgesi(yontem.sds(ss, zemin), zemin)
        kusur = {"yumusakKat": b["yumusak_kat"], "kisaKolon": b["kisa_kolon"]}
        return yontem.puan(b["kat"], bolge, perde=perde, kusurlar=kusur)

    auc_sonuc = {
        f"ryteie_{ad}": round(auc([-pp(b, o, b["perde_alan"] > 0) for b in kapsam], hedef), 3)
        for ad, o in ORANLAR.items()
    }
    auc_sonuc |= {
        "kat_sayisi": round(auc([b["kat"] for b in kapsam], hedef), 3),
        "yas": round(auc([b["yas"] for b in kapsam], hedef), 3),
        "yumusak_kat": round(auc([float(b["yumusak_kat"]) for b in kapsam], hedef), 3),
        "hassan_sozen": round(auc([-b["pi"] for b in kapsam], hedef), 3),
    }

    def bant_oranlari(perde_kurali) -> list[dict]:
        puanlar = [pp(b, ORANLAR["kotu"], perde_kurali(b)) for b in kapsam]
        sonuc = []
        for ad, alt, ust in BANTLAR:
            icinde = [h for p, h in zip(puanlar, hedef, strict=True) if alt <= p < ust]
            sonuc.append(
                {
                    "bant": ad,
                    "n": len(icinde),
                    "agirHasar": sum(icinde),
                    "oran": round(sum(icinde) / len(icinde), 3) if icinde else None,
                }
            )
        return sonuc

    ozet = {
        "veri": ACI133_KUNYE,
        "nToplam": len(tum),
        "nKapsam": len(kapsam),
        "hedef": "ağır hasar ve üstü (3 sınıflı hasarda S)",
        "pozitif": sum(hedef),
        "kritikHasarPozitif": sum(b["kritik"] == "Y" for b in kapsam),
        "katDagilimi": {str(k): sum(b["kat"] == k for b in kapsam) for k in range(1, 8)},
        "auc": auc_sonuc,
        "bantlar": {
            "perdeVerisiyle": bant_oranlari(lambda b: b["perde_alan"] > 0),
            "bacVarsayimiyla": bant_oranlari(lambda b: False),
        },
        "bantTehlikeVarsayimi": "kötü uç: SS_ESHM20 / 0,72 (ürünün bant kuralı)",
        "olculemeyenParametreler": OLCULEMEYEN,
        "yanlilik": "Örneklem hasar gözlemi için seçilmiş binalar; ağır hasar oranı genel stoğu temsil etmez. "
        "AUC sıralama gücüdür, mutlak isabet değil.",
    }
    return ozet, {"aci133": xlsx_meta, "eshm20": izgara_meta}
